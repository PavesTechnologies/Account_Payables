"""Finance dashboard + AP reports (Business_Layer/services/ap_reporting_service.py) with a fake DAO:
amount rules (TDS, partial payments, scheduled reservations - each invoice counted once), due-date
buckets, per-currency totals, permissions, the statutory TDS calendar and the Excel/PDF exports."""
import datetime
import io
from decimal import Decimal
from types import SimpleNamespace

import fitz
import pytest
from openpyxl import load_workbook

from Backend.Business_Layer.services import ap_reporting_service as svc_module
from Backend.Business_Layer.services.ap_reporting_service import APReportingService
from Backend.Business_Layer.utils.report_export import to_pdf, to_xlsx
from Backend.Business_Layer.utils.statutory_calendar import (
    financial_quarter,
    tds_deposit_due_date,
    tds_statement_due_date,
)

D = datetime.date
TODAY = D(2026, 10, 9)
FINANCE = {"permissions": ["PAYMENT_VIEW"]}
FINANCE_TDS = {"permissions": ["PAYMENT_VIEW", "TDS_TRACKING_VIEW"]}


def _row(invoice_id, status, net, due, paid="0", tds=None, term=None, tracking=None, currency=("INR", "₹"),
         vendor="Vendor A", vendor_id=1, invoice_date=D(2026, 8, 1)):
    invoice = SimpleNamespace(invoice_id=invoice_id, invoice_number=f"INV-{invoice_id}", vendor_id=vendor_id,
                              invoice_type="NON_PO", invoice_date=invoice_date, due_date=due,
                              net_amount=Decimal(net), amount_paid=Decimal(paid))
    return (invoice, vendor, status, status.title(), currency[0], currency[1], tds, term, tracking,
            "Finance", "Audit", "PROFESSIONAL_SERVICE" if tds else None, "194J" if tds else None)


def _tds(amount, status="VERIFIED"):
    return SimpleNamespace(tds_applicable=True, tds_amount=Decimal(amount), determination_status=status,
                           taxable_base=Decimal(amount) * 10, tds_rate=Decimal("10"))


def _term(status="COMPLIANT", effective=True, reason=None, statutory=None):
    return SimpleNamespace(validation_status=status, reason_code=reason,
                           effective_due_date=D(2026, 1, 1) if effective else None, statutory_due_date=statutory,
                           invoice_terms_text="Net 30", invoice_term_days=30, reference_source="VENDOR_MASTER",
                           applied_term_days=30, suggested_term_days=None)


class _FakeDAO:
    def __init__(self, rows, reserved=None, cleared=None):
        self.rows, self.reserved, self.cleared = rows, reserved or {}, cleared or []

    def invoice_rows(self, statuses=None, **filters):
        rows = [r for r in self.rows if not statuses or r[2] in statuses]
        if filters.get("tds_only"):
            rows = [r for r in rows if r[6] is not None]
        return rows

    def reserved_by_invoice(self, ids):
        return {i: Decimal(v) for i, v in self.reserved.items()}

    def cleared_payment_rows(self, paid_from=None, paid_to=None, vendor_id=None, limit=None):
        return self.cleared[:limit] if limit else self.cleared


@pytest.fixture
def make(monkeypatch):
    monkeypatch.setattr(svc_module, "configured_deposit_days", lambda db: (7, 30))
    monkeypatch.setattr(svc_module.PaymentTermDAO, "list_expiring_agreements", lambda self, today, days: [])

    def factory(rows, **kw):
        service = APReportingService(db=None, today=TODAY)
        service.dao = _FakeDAO(rows, **kw)
        return service
    return factory


def _amount(entries, code="INR"):
    return next((e["amount"] for e in entries if e["currency_code"] == code), Decimal("0.00"))


def test_outstanding_counts_tds_partial_payment_and_reservation_once(make):
    service = make([_row(1, "PARTIALLY_PAID", "11800", D(2026, 10, 20), paid="5000", tds=_tds("1000"))],
                   reserved={1: "800"})
    [figure] = service.invoice_figures()
    assert figure.net_payable == Decimal("10800.00")      # 11800 - 1000 TDS
    assert figure.outstanding == Decimal("5000.00")       # - 5000 paid - 800 scheduled
    card = {c["key"]: c for c in service.finance_dashboard(FINANCE)["action_cards"]}
    assert _amount(card["ready_for_payment"]["amounts"]) == Decimal("5000.00")
    assert _amount(card["partially_paid"]["amounts"]) == Decimal("5000.00")
    assert card["ready_for_payment"]["count"] == 1


def test_determined_but_unverified_tds_still_reduces_payable_pending_tds_does_not(make):
    service = make([_row(1, "APPROVED", "1000", D(2026, 11, 1), tds=_tds("100", "DETERMINED")),
                    _row(2, "APPROVED", "1000", D(2026, 11, 1), tds=_tds("100", "PENDING"))])
    first, second = service.invoice_figures()
    assert first.outstanding == Decimal("900.00")
    assert second.outstanding == Decimal("1000.00")


def test_overdue_and_ageing_buckets(make):
    rows = [
        _row(1, "READY_FOR_PAYMENT", "100", D(2026, 10, 1)),   # 8 days overdue
        _row(2, "APPROVED", "200", D(2026, 8, 20)),            # 50 days overdue
        _row(3, "READY_FOR_PAYMENT", "300", D(2026, 10, 12)),  # due in 3 days
        _row(4, "PENDING_APPROVAL", "400", D(2026, 9, 1)),     # in approval: not in finance ageing
        _row(5, "APPROVED", "500", D(2026, 9, 1), term=_term("REVIEW_REQUIRED", effective=False)),
    ]
    dash = make(rows).finance_dashboard(FINANCE)
    cards = {c["key"]: c for c in dash["action_cards"]}
    assert _amount(cards["overdue"]["amounts"]) == Decimal("300.00")
    ready_detail = {d["label"]: d for d in cards["ready_for_payment"]["detail"]}
    assert _amount(ready_detail["Due within 7 days"]["amounts"]) == Decimal("300.00")
    assert _amount(ready_detail["Overdue"]["amounts"]) == Decimal("100.00")
    buckets = {b["label"]: _amount(b["amounts"]) for b in dash["ageing"]["buckets"]}
    assert buckets["1-15 days overdue"] == Decimal("100.00")
    assert buckets["31-60 days overdue"] == Decimal("200.00")
    assert buckets["Not yet due"] == Decimal("300.00")
    assert buckets["Due date unverified"] == Decimal("500.00")


def test_upcoming_months_split_approved_and_pipeline_and_unverified(make):
    rows = [
        _row(1, "READY_FOR_PAYMENT", "100", D(2026, 10, 20)),
        _row(2, "PENDING_APPROVAL", "200", D(2026, 11, 5)),
        _row(3, "APPROVED", "300", D(2026, 12, 31)),
        _row(4, "APPROVED", "400", D(2027, 2, 1)),
        _row(5, "APPROVED", "500", D(2026, 9, 1)),
        _row(6, "OCR_REVIEWED", "600", D(2026, 10, 9), term=_term("REVIEW_REQUIRED", effective=False)),
        _row(7, "DISPUTED", "700", D(2026, 10, 15)),
    ]
    dash = make(rows).finance_dashboard(FINANCE)
    by_label = {b["label"]: b for b in dash["upcoming"]["buckets"]}
    assert _amount(by_label["Oct 2026"]["approved"]) == Decimal("100.00")
    assert _amount(by_label["Nov 2026"]["pipeline"]) == Decimal("200.00")
    assert _amount(by_label["Dec 2026"]["approved"]) == Decimal("300.00")
    assert _amount(by_label["After Dec 2026"]["approved"]) == Decimal("400.00")
    assert _amount(by_label["Overdue"]["approved"]) == Decimal("500.00")
    assert _amount(by_label["Due date unverified"]["pipeline"]) == Decimal("600.00")
    assert sum(b["count"] for b in dash["upcoming"]["buckets"]) == 6  # disputed excluded
    cards = {c["key"]: c for c in dash["action_cards"]}
    assert _amount(cards["on_hold"]["amounts"]) == Decimal("700.00")
    assert "not a guaranteed cash-flow forecast" in dash["forecast_note"]


def test_money_never_summed_across_currencies(make):
    rows = [_row(1, "READY_FOR_PAYMENT", "100", D(2026, 11, 1)),
            _row(2, "READY_FOR_PAYMENT", "50", D(2026, 11, 1), currency=("USD", "$"))]
    card = make(rows).finance_dashboard(FINANCE)["action_cards"][0]
    assert {(a["currency_code"], a["amount"]) for a in card["amounts"]} == {("INR", Decimal("100.00")),
                                                                              ("USD", Decimal("50.00"))}


def test_term_exceptions_and_blocked_count(make):
    rows = [_row(1, "APPROVED", "100", D(2026, 11, 1), term=_term("MISMATCH", reason="INVOICE_VS_PO"))]
    dash = make(rows).finance_dashboard(FINANCE)
    assert dash["term_exception_count"] == 1
    assert "purchase order" in dash["term_exceptions"][0]["reason_text"]
    approved = next(c for c in dash["action_cards"] if c["key"] == "approved_not_ready")
    assert approved["note"] == "1 blocked by payment-term exceptions"


def test_tds_block_only_with_tds_permission_and_deposit_due(make):
    tracking = SimpleNamespace(tracking_status="TDS_DEDUCTED", deduction_date=D(2026, 8, 25),
                               deposit_date=None, challan_number=None, filing_date=None, filing_reference=None)
    rows = [_row(1, "PAID", "1000", D(2026, 9, 1), paid="900", tds=_tds("100"), tracking=tracking)]
    rows[0][0].amount_paid = Decimal("900")
    assert make(rows).finance_dashboard(FINANCE)["tds"] is None
    items = {i["key"]: i for i in make(rows).finance_dashboard(FINANCE_TDS)["tds"]["items"]}
    assert items["deposit_overdue"]["count"] == 0  # PAID invoices are not in the open set...
    # ...but the TDS register (all invoices) flags the late deposit: due 7 Sep 2026
    report = make(rows).run_report("tds_register", FINANCE_TDS, D(2026, 4, 1), D(2026, 10, 9))
    assert report["rows"][0]["deposit_due"] == D(2026, 9, 7)
    assert report["rows"][0]["deposit_timeliness"] == "Overdue"


def test_report_permissions(make):
    service = make([])
    keys = {r["key"] for r in service.available_reports(FINANCE)}
    assert "tds_register" not in keys and "outstanding_payables" in keys
    with pytest.raises(PermissionError):
        service.run_report("tds_register", FINANCE)
    with pytest.raises(ValueError, match="Unknown report"):
        service.run_report("nope", FINANCE)
    with pytest.raises(ValueError, match="on or before"):
        service.run_report("payments_made", FINANCE, D(2026, 9, 1), D(2026, 8, 1))


def test_payments_made_report_timeliness_and_totals(make):
    cleared = [
        (10, D(2026, 9, 1), "NEFT", "UTIB123", Decimal("500"), 1, "INV-1", D(2026, 9, 5), "Vendor A", "INR", "₹", 1, D(2026, 8, 1)),
        (11, D(2026, 9, 20), "UPI", "123456789012", Decimal("250"), 2, "INV-2", D(2026, 9, 10), "Vendor B", "INR", "₹", 0, D(2026, 8, 20)),
    ]
    report = make([], cleared=cleared).run_report("payments_made", FINANCE, D(2026, 7, 1), D(2026, 9, 30))
    assert [r["timeliness"] for r in report["rows"]] == ["On time", "10 days late"]
    assert report["rows"][1]["receipt"] == "Missing"
    assert report["totals"] == [{"currency_code": "INR", "rows": 2, "amount": Decimal("750.00")}]
    assert "1 of 2" in report["notes"][0]


def test_outstanding_report_and_exports_round_trip(make):
    rows = [_row(1, "READY_FOR_PAYMENT", "1180", D(2026, 10, 1), tds=_tds("100")),
            _row(2, "APPROVED", "500", D(2026, 12, 1), vendor="Vendor B", vendor_id=2)]
    report = make(rows).run_report("outstanding_payables", FINANCE)
    assert [r["due_status"] for r in report["rows"]] == ["Overdue 8 days", "Due in 53 days"]
    assert report["totals"][0]["outstanding"] == Decimal("1580.00")

    sheet = load_workbook(io.BytesIO(to_xlsx(report))).active
    values = [[c.value for c in row] for row in sheet.iter_rows()]
    assert values[0][0] == report["title"]
    header = values[3]
    assert header[-1] == "Outstanding"
    assert values[4][-1] == 1080.0
    assert any(row[0] and str(row[0]).startswith("Total (INR)") and row[-1] == 1580.0 for row in values)

    pdf = fitz.open("pdf", to_pdf(report))
    text = pdf[0].get_text()
    assert "INV-1" in text and "1,580.00" in text and "Page 1 of 1" in text


def test_ageing_report_by_vendor(make):
    rows = [_row(1, "APPROVED", "100", D(2026, 9, 1)), _row(2, "APPROVED", "50", D(2026, 10, 30)),
            _row(3, "READY_FOR_PAYMENT", "70", D(2026, 5, 1), vendor="Vendor B", vendor_id=2)]
    report = make(rows).run_report("ageing", FINANCE)
    by_vendor = {r["vendor_name"]: r for r in report["rows"]}
    assert by_vendor["Vendor A"]["b_31-60"] == Decimal("100.00")
    assert by_vendor["Vendor A"]["not_due"] == Decimal("50.00")
    assert by_vendor["Vendor B"]["b_90+"] == Decimal("70.00")
    assert by_vendor["Vendor A"]["total"] == Decimal("150.00")


# ---------------------------------------------------------------------------
# Statutory calendar
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("deducted, due", [
    (D(2026, 8, 25), D(2026, 9, 7)),
    (D(2026, 12, 3), D(2027, 1, 7)),
    (D(2027, 3, 31), D(2027, 4, 30)),   # March deductions: 30 April
    (D(2027, 1, 31), D(2027, 2, 7)),
])
def test_tds_deposit_due(deducted, due):
    assert tds_deposit_due_date(deducted) == due


@pytest.mark.parametrize("deducted, due", [
    (D(2026, 4, 10), D(2026, 7, 31)),
    (D(2026, 9, 30), D(2026, 10, 31)),
    (D(2026, 12, 1), D(2027, 1, 31)),
    (D(2027, 2, 15), D(2027, 5, 31)),
])
def test_quarterly_statement_due(deducted, due):
    assert tds_statement_due_date(deducted) == due


def test_financial_quarter():
    assert financial_quarter(D(2026, 8, 15)) == ("2026-27", 2)
    assert financial_quarter(D(2027, 3, 31)) == ("2026-27", 4)
    assert tds_deposit_due_date(None) is None and tds_statement_due_date(None) is None


# ---------------------------------------------------------------------------
# Headline KPIs, paid trend and the Reports period summary
# ---------------------------------------------------------------------------

def test_finance_kpis_and_paid_trend(make):
    rows = [_row(1, "READY_FOR_PAYMENT", "100", D(2026, 10, 1)),          # overdue
            _row(2, "APPROVED", "200", D(2026, 10, 12)),                  # due within 7 days
            _row(3, "APPROVED", "50", D(2026, 11, 30), term=_term("MISMATCH", reason="INVOICE_VS_PO"))]
    cleared = [(10, D(2026, 10, 2), "NEFT", "R1", Decimal("300"), 9, "INV-9", D(2026, 10, 5), "V", "INR", "₹", 1, D(2026, 9, 1)),
               (11, D(2026, 9, 2), "NEFT", "R2", Decimal("40"), 8, "INV-8", D(2026, 9, 5), "V", "INR", "₹", 1, D(2026, 8, 1)),
               (12, D(2026, 10, 3), "NEFT", "R3", Decimal("9"), 7, "INV-7", D(2026, 10, 5), "V", "USD", "$", 1, D(2026, 9, 1))]
    dash = make(rows, cleared=cleared).finance_dashboard(FINANCE)
    kpis = {k["key"]: k for k in dash["kpis"]}
    assert kpis["ready_to_pay"]["value"] == Decimal("100.00")
    assert kpis["due_this_week"]["value"] == Decimal("200.00")
    assert kpis["overdue"]["value"] == Decimal("100.00")
    assert kpis["paid_this_month"]["value"] == Decimal("300.00")      # USD payment not mixed in
    assert kpis["needs_action"]["value"] == 3                           # 2 approved-not-ready + 1 term exception
    assert [t["label"] for t in dash["paid_trend"]] == ["May", "Jun", "Jul", "Aug", "Sep", "Oct"]
    assert dash["paid_trend"][-2]["amount"] == Decimal("40.00")


def test_report_summary_period_figures(make):
    rows = [_row(1, "PAID", "1180", D(2026, 9, 1), paid="1080", tds=_tds("100"), invoice_date=D(2026, 8, 1)),
            _row(2, "APPROVED", "500", D(2026, 9, 15), invoice_date=D(2026, 9, 1), vendor="Vendor B", vendor_id=2),
            _row(3, "REJECTED", "999", D(2026, 9, 15), invoice_date=D(2026, 9, 1))]
    cleared = [(10, D(2026, 8, 31), "NEFT", "R1", Decimal("1080"), 1, "INV-1", D(2026, 9, 1), "Vendor A", "INR", "₹", 1, D(2026, 8, 1))]
    summary = make(rows, cleared=cleared).report_summary(FINANCE, D(2026, 8, 1), D(2026, 10, 9))
    kpis = {k["key"]: k["value"] for k in summary["kpis"]}
    assert kpis["invoiced"] == Decimal("1680.00")      # rejected excluded
    assert kpis["paid"] == Decimal("1080.00")
    assert kpis["tds"] == Decimal("100.00")
    assert kpis["outstanding"] == Decimal("500.00")
    assert kpis["overdue"] == Decimal("500.00")
    assert kpis["on_time"] == 100.0 and kpis["days_to_pay"] == 30.0 and kpis["vendors"] == 2
    by_month = {m["period"]: m for m in summary["monthly"]}
    assert by_month["2026-08"]["invoiced"] == Decimal("1180.00") and by_month["2026-08"]["paid"] == Decimal("1080.00")
    assert [v["label"] for v in summary["by_vendor"]] == ["Vendor A", "Vendor B"]
    with pytest.raises(PermissionError):
        make([]).report_summary({"permissions": ["INVOICE_VIEW"]})


def test_management_report_access(make):
    service = make([])
    viewer = {"permissions": ["AP_MANAGEMENT_REPORTS_VIEW"]}
    keys = {r["key"]: r for r in service.available_reports(viewer)}
    assert "tds_register" in keys and keys["tds_register"]["can_export"] is False
    service.run_report("outstanding_payables", viewer)  # read is allowed
    exporter = {"permissions": ["AP_MANAGEMENT_REPORTS_VIEW", "AP_MANAGEMENT_REPORTS_EXPORT"]}
    assert service.can_export(exporter, "tds_register") is True
    assert service.can_export({"permissions": ["PAYMENT_VIEW"]}, "payments_made") is True
