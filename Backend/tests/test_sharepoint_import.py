"""Daily SharePoint pickup (6.7): dry run, import + auto-record + move to Processed with a results
sheet, once per file version, rejected sheets, and the on/off switch. Fake SharePoint + importer."""
import asyncio
import datetime
from types import SimpleNamespace

import pytest

from Backend.Business_Layer.services import sharepoint_import_service as sp
from Backend.Business_Layer.utils import bulk_import_workbook as wbk
from Backend.Business_Layer.utils import graph_drive_client as gdc
from Backend.Data_Access_Layer.models.audit import AuditLog

TODAY = datetime.date(2026, 10, 12)


def _payment_sheet(receipt="advice-7.pdf"):
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(wbk.build_template("PAYMENTS", [{"invoice_id": 7, "invoice_number": "INV-7"}])))
    ws = wb.active
    ws["E2"], ws["F2"], ws["G2"], ws["H2"], ws["I2"] = "NEFT", datetime.date(2026, 10, 8), 1180, "SBIN426281234567", receipt
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class FakeDrive:
    def __init__(self, inbox):
        self.inbox = inbox  # {"Payments/Inbox": [items]}
        self.files = {}
        self.moved, self.uploaded, self.folders = [], [], []

    def children(self, folder):
        return self.inbox.get(folder, [])

    def download(self, item_id):
        return self.files[item_id]

    def ensure_folder(self, path):
        self.folders.append(path)
        return f"folder:{path}"

    def move(self, item_id, folder_id, new_name=None):
        self.moved.append((item_id, folder_id))

    def upload(self, folder_id, name, content):
        self.uploaded.append((folder_id, name))


def _item(item_id, name, etag="v1"):
    return {"id": item_id, "name": name, "eTag": etag, "file": {}, "lastModifiedBy": {"user": {"displayName": "Priya (Finance)"}}}


@pytest.fixture
def world(monkeypatch):
    state = SimpleNamespace(jobs={}, existing=set(), validated=[], created=[], reject=None)

    class Db:
        def query(self, model):
            class Q:
                def filter(self, *a):
                    return self

                def first(self_inner):
                    return object() if state.check_ref in state.existing else None
            return Q()

        def rollback(self):
            pass

        def close(self):
            pass

    class Service:
        def __init__(self, db):
            pass

        def create_job(self, kind, workbook, documents, user_id, user_name, source, source_reference, auto_record):
            if state.reject:
                raise ValueError(state.reject)
            state.created.append({"kind": kind, "sheet": workbook[0], "docs": [d[0] for d in documents], "auto_record": auto_record,
                                  "source": source, "by": user_name})
            state.existing.add(source_reference)
            return 21

        def job_detail(self, job_id):
            return {"status": "COMPLETED", "counts": {"recorded": 1, "needs_review": 0, "error": 0}}

        def results_workbook(self, job_id):
            return "pay - results.xlsx", b"xlsx"

    async def validate(job_id, actor):
        state.validated.append((job_id, actor))

    def session():
        return Db()

    original_import = sp._import

    async def tracking_import(kind, folder, sheet, book, docs, client, today):
        state.check_ref = f"{sheet['id']}:{sheet.get('eTag')}"
        return await original_import(kind, folder, sheet, book, docs, client, today)

    monkeypatch.setattr(sp, "SessionLocal", session)
    monkeypatch.setattr(sp.importer, "BulkImportService", Service)
    monkeypatch.setattr(sp.importer, "validate_job", validate)
    monkeypatch.setattr(sp, "_import", tracking_import)
    return state


def _drive():
    drive = FakeDrive({"Payments/Inbox": [_item("s1", "pay.xlsx"), _item("r1", "advice-7.pdf"), _item("x", "unrelated.pdf")]})
    drive.files = {"s1": _payment_sheet(), "r1": b"%PDF receipt"}
    return drive


def test_dry_run_lists_without_importing(world):
    drive = _drive()
    report = asyncio.run(sp.run_once(False, client=drive, today=TODAY))
    assert report.sheets == [{"type": "PAYMENTS", "sheet": "pay.xlsx", "modified_by": "Priya (Finance)", "documents": ["advice-7.pdf"]}]
    assert world.created == [] and drive.moved == []


def test_execute_imports_records_and_moves_to_processed(world):
    drive = _drive()
    report = asyncio.run(sp.run_once(True, client=drive, today=TODAY))
    assert world.created == [{"kind": "PAYMENTS", "sheet": "pay.xlsx", "docs": ["advice-7.pdf"], "auto_record": True,
                              "source": "SHAREPOINT", "by": "SharePoint (Priya (Finance))"}]
    assert world.validated == [(21, "SHAREPOINT_IMPORT")]
    folder = "folder:Payments/Processed/2026-10-12"
    assert (folder, "pay - results.xlsx") in drive.uploaded
    assert ("s1", folder) in drive.moved and ("r1", folder) in drive.moved and ("x", folder) not in drive.moved
    assert report.sheets[0]["job_id"] == 21

    # the same file version again -> skipped
    again = asyncio.run(sp.run_once(True, client=_drive(), today=TODAY))
    assert again.sheets[0]["skipped"] == "already imported" and len(world.created) == 1


def test_unreadable_sheet_is_moved_out_with_a_note(world):
    world.reject = "Missing column(s): Amount Paid. Download a fresh template."
    drive = _drive()
    report = asyncio.run(sp.run_once(True, client=drive, today=TODAY))
    assert "Missing column" in report.sheets[0]["rejected"]
    assert any(name.endswith("NOT IMPORTED.txt") for _, name in drive.uploaded) and ("s1", "folder:Payments/Processed/2026-10-12") in drive.moved


def test_documents_inside_a_zip_are_picked_up(world):
    drive = FakeDrive({"Payments/Inbox": [_item("s1", "pay.xlsx"), _item("z1", "receipts.zip")]})
    drive.files = {"s1": _payment_sheet(), "z1": b"PK zip"}
    asyncio.run(sp.run_once(True, client=drive, today=TODAY))
    assert world.created[0]["docs"] == ["receipts.zip"]


# ======================================================================
# Switch + lock
# ======================================================================
class _CfgDb:
    def __init__(self):
        self.rows, self.audit = {}, []

    def add(self, obj):
        (self.audit.append(obj) if isinstance(obj, AuditLog) else self.rows.__setitem__(obj.config_key, obj))

    def commit(self):
        pass

    def close(self):
        pass


@pytest.fixture
def cfg(monkeypatch):
    db = _CfgDb()

    class MasterDAO:
        def __init__(self, d):
            pass

        def get_system_config_by_key(self, key):
            return db.rows.get(key)
    monkeypatch.setattr(sp, "MasterDAO", MasterDAO)
    return db


def test_switch_off_by_default_and_audited(cfg):
    settings = sp.SharePointImportSettings(cfg)
    assert settings.is_enabled() is False
    settings.set_enabled(True, "u-1", "Ravi (Finance Manager)")
    assert settings.is_enabled() and cfg.audit[-1].action == "SHAREPOINT_IMPORT_ENABLED"
    assert settings.status()["updated_by"] == "Ravi (Finance Manager)"


def test_execute_does_nothing_while_off(cfg, monkeypatch):
    called = []

    async def fake_run(execute, client=None, today=None):
        called.append(execute)
        return sp.PickupReport()
    monkeypatch.setattr(sp, "SessionLocal", lambda: cfg)
    monkeypatch.setattr(sp, "try_lock", lambda db: True)
    monkeypatch.setattr(sp, "unlock", lambda db: None)
    monkeypatch.setattr(sp, "run_once", fake_run)
    assert asyncio.run(sp.run_locked(True)).disabled is True and called == []
    asyncio.run(sp.run_locked(False))  # dry run ignores the switch
    sp.SharePointImportSettings(cfg).set_enabled(True, "u-1")
    asyncio.run(sp.run_locked(True))
    assert called == [False, True] and '"status":"ok"' in cfg.rows["BULK_IMPORT_SHAREPOINT_LAST_RUN"].config_value


def test_missing_inbox_folder_is_empty(monkeypatch):
    env = {"TENANT_ID": "t", "CLIENT_ID": "c", "CLIENT_SECRET": "s", "SHAREPOINT_SITE": "x.sharepoint.com:/sites/AP"}
    monkeypatch.setattr(gdc, "get_env_var", lambda k, d=None: env.get(k, d))
    import Backend.Business_Layer.utils.graph_mail_client as gmc
    monkeypatch.setattr(gmc, "get_env_var", lambda k, d=None: env.get(k, d))
    client = gdc.GraphDriveClient(session=SimpleNamespace())
    monkeypatch.setattr(client, "drive_id", lambda: "d1")

    def not_found(url, params=None, raw=False):
        raise gmc.GraphMailError("Microsoft Graph request failed (HTTP 404, itemNotFound).")
    monkeypatch.setattr(client, "_get", not_found)
    assert client.children("Payments/Inbox") == []
    with pytest.raises(gdc.GraphDriveError, match="SHAREPOINT_SITE"):
        env.pop("SHAREPOINT_SITE")
        gdc.GraphDriveClient(session=SimpleNamespace())


# ======================================================================
# Local folder (testing without SharePoint)
# ======================================================================
def test_local_folder_client_and_full_pickup(world, tmp_path):
    from Backend.Business_Layer.utils.local_drive_client import LocalDriveClient
    client = LocalDriveClient(str(tmp_path))
    inbox = tmp_path / "Payments" / "Inbox"
    inbox.mkdir(parents=True)
    (inbox / "pay.xlsx").write_bytes(_payment_sheet())
    (inbox / "advice-7.pdf").write_bytes(b"%PDF receipt")
    (inbox / "~$pay.xlsx").write_bytes(b"excel lock file")
    items = client.children("Payments/Inbox")
    assert [i["name"] for i in items] == ["advice-7.pdf", "pay.xlsx"] and client.children("TDS Filings/Inbox") == []

    report = asyncio.run(sp.run_once(True, client=client, today=TODAY))
    assert report.sheets[0]["job_id"] == 21 and world.created[0]["by"] == "Local folder"
    processed = tmp_path / "Payments" / "Processed" / "2026-10-12"
    assert sorted(p.name for p in processed.iterdir()) == ["advice-7.pdf", "pay - results.xlsx", "pay.xlsx"]
    assert not (inbox / "pay.xlsx").exists()
    with pytest.raises(ValueError, match="outside"):
        client.download("local:" + str(tmp_path.parent / "secret.txt"))
