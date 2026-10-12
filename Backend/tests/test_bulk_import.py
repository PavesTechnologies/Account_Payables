"""Bulk import (6.7): Excel template / parsing / results, document expansion, row checks against
the documents, recording through the existing services, and per-type permissions. No DB / S3 /
Textract."""
import datetime
import io
import zipfile
from decimal import Decimal
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from starlette.middleware.base import BaseHTTPMiddleware

from Backend.API_Layer.routes import bulk_import_route as route
from Backend.Business_Layer.services import bulk_import_service as svc
from Backend.Business_Layer.utils import bulk_import_workbook as wbk

PDF = b"%PDF-1.7 receipt"
TODAY = datetime.date(2026, 10, 12)


# ======================================================================
# Workbook
# ======================================================================
def test_template_round_trip_and_info_columns_are_ignored():
    content = wbk.build_template("PAYMENTS", [{"invoice_id": 7, "invoice_number": "INV-7", "vendor_name": "Acme",
                                               "remaining_amount": Decimal("1180.00"), "amount": Decimal("1180.00")}])
    wb = load_workbook(io.BytesIO(content))
    ws = wb["Payments"]
    assert ws["B1"].value == "Invoice Number *" and "Instructions" in wb.sheetnames
    assert ws.data_validations.dataValidation[0].formula1.startswith('"NEFT,RTGS')
    ws["E2"], ws["F2"], ws["H2"], ws["I2"] = "neft", datetime.date(2026, 10, 8), "SBIN426281234567", "advice-7.pdf"
    out = io.BytesIO()
    wb.save(out)
    rows = wbk.parse_workbook("PAYMENTS", out.getvalue())
    assert len(rows) == 1 and rows[0].problems == []
    assert rows[0].data == {"invoice_id": 7, "invoice_number": "INV-7", "payment_date": datetime.date(2026, 10, 8),
                            "amount": Decimal("1180.00"), "payment_mode": "NEFT", "reference_number": "SBIN426281234567",
                            "receipt_file": "advice-7.pdf", "remarks": None}
    assert "vendor_name" not in rows[0].data  # grey columns are never read


def _sheet(headers, *rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_parse_reports_row_problems_and_rejects_bad_workbooks():
    headers = ["Invoice ID", "Invoice Number *", "Payment Date *", "Amount Paid *", "Payment Mode *", "UTR / Reference *", "Receipt File Name *"]
    rows = wbk.parse_workbook("PAYMENTS", _sheet(headers, [1, "INV-1", "31-02-2026", "abc", "CASH", 123456789012, None], [None] * 7))
    assert len(rows) == 1  # blank row skipped
    problems = " | ".join(rows[0].problems)
    assert "Payment Date" in problems and "Amount Paid" in problems and "Payment Mode" in problems and "Receipt File Name is required" in problems
    assert rows[0].data["reference_number"] == "123456789012"  # numeric UTR typed into Excel stays text
    with pytest.raises(wbk.WorkbookRejected, match="Missing column"):
        wbk.parse_workbook("PAYMENTS", _sheet(["Invoice Number"], ["INV-1"]))
    with pytest.raises(wbk.WorkbookRejected, match="not a readable"):
        wbk.parse_workbook("PAYMENTS", b"not excel")


def test_group_keys():
    assert wbk.group_key("PAYMENTS", {"reference_number": " sbin 4262 "}) == "UTR:SBIN4262"
    assert wbk.group_key("TDS_CHALLANS", {"challan_serial_no": "123", "bsr_code": "0510308", "deposit_date": datetime.date(2026, 10, 7)}) \
        == "CHALLAN:0510308:2026-10-07:00123"
    assert wbk.group_key("TDS_FILINGS", {"acknowledgement_no": "1234 5678"}) == "FILING:12345678"
    assert wbk.group_key("TDS_CHALLANS", {"bsr_code": "0510308"}) is None


def test_results_workbook():
    content = wbk.build_results("PAYMENTS", [{"data": {"invoice_number": "INV-1", "amount": "100.00"}, "status": "RECORDED",
                                              "messages": [{"severity": "info", "message": "ok"}], "result_reference": "PAYMENT-9"}])
    ws = load_workbook(io.BytesIO(content)).active
    headers = [c.value for c in ws[1]]
    assert headers[-3:] == ["Import Status", "Messages", "Recorded Reference"]
    assert [c.value for c in ws[2]][-3:] == ["RECORDED", "ok", "PAYMENT-9"]


def test_expand_documents_loose_and_zip():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("receipts/advice-1.pdf", PDF)
        zf.writestr("__MACOSX/._advice-1.pdf", b"x")
    out = svc.expand_documents([("advice-2.pdf", "application/pdf", PDF + b"2"), ("all.zip", "application/zip", buf.getvalue())])
    assert [c.file_name for c in out] == ["advice-2.pdf", "advice-1.pdf"] and all(c.problem is None for c in out)
    with pytest.raises(ValueError, match="not a valid ZIP"):
        svc.expand_documents([("bad.zip", None, b"nope")])


# ======================================================================
# Payment checks
# ======================================================================
def _row(row_id, data, doc="docs/a.pdf", group=None):
    return SimpleNamespace(row_id=row_id, row_number=row_id + 1, data=data, status="PENDING", messages=None, extracted=None,
                           invoice_id=None, document_path=doc, document_name="a.pdf",
                           group_key=group or wbk.group_key("PAYMENTS", {"reference_number": data.get("reference_number")}))


def _pay(invoice_id=7, amount="1180.00", ref="SBIN426281234567", date="2026-10-08", **kw):
    return {"invoice_id": invoice_id, "invoice_number": f"INV-{invoice_id}", "payment_date": date, "amount": amount,
            "payment_mode": "NEFT", "reference_number": ref, "receipt_file": "a.pdf", **kw}


def _receipt(amount="1180.00", ref="SBIN426281234567", date="2026-10-08", status="SUCCESS", beneficiary="ACME SUPPLIES PVT LTD"):
    return {"fields": {"amount": {"value": amount}, "reference_number": {"value": ref}, "payment_date": {"value": date}},
            "transaction_status": status, "beneficiary_name": beneficiary}


@pytest.fixture
def payments(monkeypatch):
    invoices = {7: {"invoice_number": "INV-7", "vendor_name": "Acme Supplies Pvt Ltd", "remaining_amount": "1180.00", "can_record_payment": True, "status_code": "READY_FOR_PAYMENT"},
                8: {"invoice_number": "INV-8", "vendor_name": "Acme Supplies Pvt Ltd", "remaining_amount": "500.00", "can_record_payment": True, "status_code": "READY_FOR_PAYMENT"},
                9: {"invoice_number": "INV-9", "vendor_name": "Zeta", "remaining_amount": "100.00", "can_record_payment": False, "status_name": "Paid"}}

    class Tracking:
        def __init__(self, db):
            pass

        def get_invoice_payments(self, invoice_id):
            if invoice_id not in invoices:
                raise LookupError("not found")
            return invoices[invoice_id]
    monkeypatch.setattr(svc, "PaymentTrackingService", Tracking)
    return svc.BulkImportService(None, today=TODAY)


def _evaluate(service, rows, extracted):
    service._evaluate_payments(rows, extracted)
    return [(r.status, [m["message"] for m in (r.messages or [])]) for r in rows]


def test_payment_confirmed_by_receipt_is_ready(payments):
    [(status, msgs)] = _evaluate(payments, [_row(1, _pay())], {"docs/a.pdf": _receipt()})
    assert status == "READY" and msgs == []


def test_shared_utr_rows_must_add_up_to_the_receipt(payments):
    rows = [_row(1, _pay(7, "1180.00")), _row(2, _pay(8, "500.00"))]
    assert {s for s, _ in _evaluate(payments, rows, {"docs/a.pdf": _receipt("1680.00")})} == {"READY"}
    rows = [_row(1, _pay(7, "1180.00")), _row(2, _pay(8, "400.00"))]
    out = _evaluate(payments, rows, {"docs/a.pdf": _receipt("1680.00")})
    assert out[0][0] == "NEEDS_REVIEW" and "rows with this UTR add up to 1,580.00" in out[0][1][0]


@pytest.mark.parametrize("receipt,expected", [
    (_receipt(ref="SBIN000000000000"), "differs from the sheet"),
    (_receipt(date="2026-10-09"), "Receipt date 09-10-2026 differs"),
    (_receipt(amount="1000.00"), "Receipt amount 1,000.00"),
    (_receipt(status="PENDING"), "pending"),
    (_receipt(beneficiary="Zeta Traders"), "does not look like the vendor"),
    (None, "could not be read"),
])
def test_receipt_mismatches_need_review(payments, receipt, expected):
    [(status, msgs)] = _evaluate(payments, [_row(1, _pay())], {"docs/a.pdf": receipt} if receipt else {})
    assert status == "NEEDS_REVIEW" and any(expected in m for m in msgs)


@pytest.mark.parametrize("data,receipt,expected", [
    (_pay(9, "100.00"), _receipt("100.00"), "not ready for payment"),
    (_pay(7, "2000.00"), _receipt("2000.00"), "exceed the remaining payable"),
    (_pay(42), _receipt(), "not found"),
    (_pay(ref="SHORT1"), _receipt(ref="SHORT1"), "not a valid NEFT reference"),
    (_pay(date="2026-12-01"), _receipt(date="2026-12-01"), "in the future"),
    (_pay(), _receipt(status="FAILED"), "failed or was reversed"),
])
def test_payment_errors(payments, data, receipt, expected):
    [(status, msgs)] = _evaluate(payments, [_row(1, data)], {"docs/a.pdf": receipt})
    assert status == "ERROR" and any(expected in m for m in msgs)


def test_two_rows_for_one_invoice_cannot_exceed_its_balance(payments):
    rows = [_row(1, _pay(7, "700.00", ref="SBIN426281234567")), _row(2, _pay(7, "700.00", ref="SBIN426281234568"), doc="docs/b.pdf")]
    out = _evaluate(payments, rows, {"docs/a.pdf": _receipt("700.00"), "docs/b.pdf": _receipt("700.00", ref="SBIN426281234568")})
    assert out[0][0] == "READY" and out[1][0] == "ERROR"


# ======================================================================
# Challan / filing groups
# ======================================================================
def test_challan_group_uses_validate_challan_and_the_document(monkeypatch):
    seen = {}

    class Challans:
        def __init__(self, db, today=None):
            pass

        def validate_challan(self, header, allocations, has_document):
            seen.update(header=header, allocations=allocations)
            return {"errors": [], "warnings": [{"message": "Deposited after the due date"}],
                    "normalized": {"challan_serial_no": "00123", "bsr_code": "0510308", "deposit_date": datetime.date(2026, 10, 7),
                                   "tax_amount": Decimal("12000.00")}}
    monkeypatch.setattr(svc, "TdsChallanService", Challans)
    common = {"challan_serial_no": "123", "bsr_code": "0510308", "deposit_date": "2026-10-07", "tax_amount": "12000.00"}
    group = [_row(1, {**common, "invoice_id": 1}, group="G"), _row(2, {**common, "invoice_id": 2, "allocated_tds_amount": "7000.00"}, group="G")]
    doc = {"fields": {"challan_serial_no": {"value": "00123"}, "bsr_code": {"value": "0510308"}, "deposit_date": {"value": "2026-10-07"},
                      "tax_amount": {"value": "12000.00"}}}
    service = svc.BulkImportService(None, today=TODAY)
    service._evaluate_challan(group, {"docs/a.pdf": doc})
    assert [r.status for r in group] == ["READY", "READY"] and group[0].messages[0]["severity"] == "info"
    assert seen["allocations"] == [{"invoice_id": 1}, {"invoice_id": 2, "allocated_tds_amount": "7000.00"}]
    doc["fields"]["tax_amount"]["value"] = "11000.00"
    group = [_row(1, {**common, "invoice_id": 1}, group="G")]
    service._evaluate_challan(group, {"docs/a.pdf": doc})
    assert group[0].status == "NEEDS_REVIEW" and any("tax amount 11000.00" in m["message"] for m in group[0].messages)


def test_filing_group_errors_from_validate_filing(monkeypatch):
    class Challans:
        def __init__(self, db, today=None):
            pass

        def validate_filing(self, header, ids, has_document):
            return {"errors": [{"message": "INV-2: TDS must be deposited first"}], "warnings": [],
                    "normalized": {"acknowledgement_no": "123456789012345", "filing_date": datetime.date(2026, 10, 9)}}
    monkeypatch.setattr(svc, "TdsChallanService", Challans)
    common = {"form_type": "26Q", "financial_year": "2026-27", "quarter": 2, "acknowledgement_no": "123456789012345", "filing_date": "2026-10-09"}
    group = [_row(1, {**common, "invoice_id": 1}, group="F"), _row(2, {**common, "invoice_id": 2, "quarter": 3}, group="F")]
    svc.BulkImportService(None, today=TODAY)._evaluate_filing(group, {"docs/a.pdf": {"fields": {"acknowledgement_no": {"value": "123456789012345"}, "filing_date": {"value": "2026-10-09"}}}})
    msgs = [m["message"] for m in group[0].messages]
    assert group[0].status == "ERROR" and "INV-2: TDS must be deposited first" in msgs and any("disagree on quarter" in m for m in msgs)


# ======================================================================
# Recording
# ======================================================================
class _Db:
    def __init__(self):
        self.commits = 0

    def commit(self):
        self.commits += 1

    def rollback(self):
        pass

    def merge(self, obj):
        return obj


def test_record_payments_through_the_existing_service(monkeypatch):
    recorded, attached = [], []

    class Tracking:
        def __init__(self, db):
            pass

        def record_payment(self, invoice_id, data, user_id):
            if invoice_id == 8:
                raise ValueError("Payment Amount 600 exceeds the remaining payable amount 500")
            recorded.append((invoice_id, data.entry_mode, str(data.amount)))
            return {"recorded_payment_id": 100 + invoice_id}

        def upload_document(self, payment_id, name, content, content_type, document_type, user_id):
            attached.append((payment_id, name, document_type, content_type))
    monkeypatch.setattr(svc, "PaymentTrackingService", Tracking)
    monkeypatch.setattr(svc, "get_object_bytes", lambda key: PDF)
    rows = [_row(1, _pay(7)), _row(2, _pay(8, "600.00")), _row(3, _pay(9))]
    rows[0].status, rows[1].status, rows[2].status = "READY", "READY", "NEEDS_REVIEW"
    for r, i in zip(rows, (7, 8, 9)):
        r.invoice_id, r.job_id = i, 5
    job = SimpleNamespace(job_id=5, kind="PAYMENTS", rows=rows, status="VALIDATED", completed_at=None, updated_at=None)
    service = svc.BulkImportService(_Db(), today=TODAY)
    monkeypatch.setattr(service, "_job", lambda job_id: job)
    summary = service.record(5, "u-1")
    assert summary == {"recorded": 1, "failed": 1}
    # the receipt is stored with its real type, so "View" opens it instead of downloading it
    assert recorded == [(7, "BULK_IMPORT", "1180.00")] and attached == [(107, "a.pdf", "RECEIPT", "application/pdf")]
    assert [r.status for r in rows] == ["RECORDED", "ERROR", "NEEDS_REVIEW"] and rows[0].result_reference == "PAYMENT-107"
    assert "exceeds the remaining payable" in rows[1].messages[-1]["message"]
    assert job.status == "NEEDS_REVIEW"
    # a person accepts the reviewed row
    service.record(5, "u-1", row_ids=[3], accept_review=True)
    assert rows[2].status == "RECORDED"


def test_challan_group_is_recorded_only_when_every_row_is_ready(monkeypatch):
    created = []

    class Challans:
        def __init__(self, db, today=None):
            pass

        def create_challan(self, header, allocations, user_id, document, source):
            created.append((len(allocations), source, document[0]))
            return {"challan_id": 3}
    monkeypatch.setattr(svc, "TdsChallanService", Challans)
    monkeypatch.setattr(svc, "get_object_bytes", lambda key: PDF)
    common = {"challan_serial_no": "123", "bsr_code": "0510308", "deposit_date": "2026-10-07", "tax_amount": "12000.00"}
    a = _row(1, {**common, "invoice_id": 1}, group="G1")
    b = _row(2, {**common, "invoice_id": 2}, group="G1")
    c = _row(3, {**common, "invoice_id": 3}, group="G2")
    a.status, b.status, c.status = "READY", "NEEDS_REVIEW", "READY"
    job = SimpleNamespace(job_id=6, kind="TDS_CHALLANS", rows=[a, b, c], status="VALIDATED", completed_at=None, updated_at=None)
    service = svc.BulkImportService(_Db(), today=TODAY)
    monkeypatch.setattr(service, "_job", lambda job_id: job)
    assert service.record(6, "u-1") == {"recorded": 1, "failed": 0}
    assert created == [(1, "BULK_IMPORT", "a.pdf")] and c.result_reference == "CHALLAN-3" and a.status == "READY"


# ======================================================================
# Routes
# ======================================================================
def _client(permissions, monkeypatch):
    class Fake:
        def __init__(self, db):
            pass

        def template(self, kind):
            return b"xlsx-" + kind.encode()

        def create_job(self, kind, workbook, docs, user_id, name):
            return 11

        def job_detail(self, job_id):
            return {"job_id": job_id}

        def _job(self, job_id):
            return SimpleNamespace(kind="TDS_CHALLANS")

        def list_jobs(self, kind, page, size):
            return {"items": [], "total": 0}
    monkeypatch.setattr(route, "BulkImportService", Fake)
    started = []

    async def fake_validate(job_id, actor):
        started.append(job_id)
    monkeypatch.setattr(route.importer, "validate_job", fake_validate)

    class _Auth(BaseHTTPMiddleware):
        async def dispatch(self, request, call_next):
            request.state.user = {"user_id": "u-1", "permissions": permissions}
            request.state.db = SimpleNamespace(rollback=lambda: None)
            return await call_next(request)
    app = FastAPI()
    app.add_middleware(_Auth)
    app.include_router(route.router, prefix="/bulk-import")
    return TestClient(app), started


def test_route_permission_per_import_type(monkeypatch):
    finance, started = _client(["PAYMENT_PROCESS"], monkeypatch)
    assert finance.get("/bulk-import/payments/template").content == b"xlsx-PAYMENTS"
    assert finance.get("/bulk-import/tds-challans/template").status_code == 403
    res = finance.post("/bulk-import/payments", files=[("workbook", ("pay.xlsx", b"PK", "application/octet-stream")),
                                                        ("documents", ("a.pdf", PDF, "application/pdf"))])
    assert res.status_code == 202 and started == [11]
    assert finance.post("/bulk-import/payments", files=[("workbook", ("pay.csv", b"x", "text/csv"))]).status_code == 400
    assert finance.get("/bulk-import/jobs/4").status_code == 403  # that job is a TDS challan import
    assert finance.get("/bulk-import/nonsense/template").status_code == 404
    viewer, _ = _client(["PAYMENT_VIEW"], monkeypatch)
    assert viewer.get("/bulk-import/payments/template").status_code == 403


def test_bulk_import_queries_compile(monkeypatch):
    from sqlalchemy.dialects import postgresql
    from sqlalchemy.orm import Query, Session

    from Backend.Data_Access_Layer import models  # noqa: F401
    compiled = []
    monkeypatch.setattr(Query, "all", lambda self: compiled.append(str(self.statement.compile(dialect=postgresql.dialect()))) or [])
    monkeypatch.setattr(Query, "count", lambda self: 0)
    service = svc.BulkImportService(Session(), today=TODAY)
    service._invoice_by_number("INV-1")
    service.list_jobs("PAYMENTS", 1, 20)
    assert len(compiled) == 2 and "bulk_import_job" in compiled[1]


def test_indian_amounts_in_exports():
    from Backend.Business_Layer.utils.report_export import INDIAN_MONEY_FORMAT, indian_amount
    assert [indian_amount(v) for v in (999, 123456, 1234567.5, 123456789.12, -45000)] == \
        ["999.00", "1,23,456.00", "12,34,567.50", "12,34,56,789.12", "-45,000.00"]
    content = wbk.build_template("PAYMENTS", [{"invoice_id": 1, "invoice_number": "INV-1", "remaining_amount": Decimal("1234567.50")}])
    ws = load_workbook(io.BytesIO(content)).active
    assert ws["D2"].number_format == INDIAN_MONEY_FORMAT


# ======================================================================
# Receipt-first: only invoice + receipt in the sheet
# ======================================================================
def _receipt_only(invoice_id=7, doc="docs/a.pdf", **kw):
    data = {"invoice_id": invoice_id, "invoice_number": f"INV-{invoice_id}", "payment_mode": "NEFT", "receipt_file": "a.pdf", **kw}
    return _row(invoice_id, data, doc=doc, group="none")


def _full_receipt(amount="1180.00", mode="NEFT"):
    r = _receipt(amount=amount)
    r["fields"]["payment_mode"] = {"value": mode}
    return r


def test_blank_fields_are_taken_from_the_receipt(payments):
    row = _receipt_only()
    [(status, msgs)] = _evaluate(payments, [row], {"docs/a.pdf": _full_receipt()})
    assert status == "READY", msgs
    assert row.data["payment_date"] == "2026-10-08" and row.data["amount"] == "1180.00"
    assert row.data["payment_mode"] == "NEFT" and row.data["reference_number"] == "SBIN426281234567"
    assert row.group_key == "UTR:SBIN426281234567"
    assert "Taken from the receipt: payment date, UTR, amount" in msgs


def test_typed_values_are_kept_and_checked(payments):
    row = _receipt_only(amount="1000.00")
    [(status, msgs)] = _evaluate(payments, [row], {"docs/a.pdf": _full_receipt("1180.00")})
    assert status == "NEEDS_REVIEW" and row.data["amount"] == "1000.00" and any("Receipt amount 1,180.00" in m for m in msgs)


def test_shared_receipt_needs_each_amount(payments):
    rows = [_receipt_only(7), _receipt_only(8)]
    out = _evaluate(payments, rows, {"docs/a.pdf": _full_receipt("1680.00")})
    assert all(s == "ERROR" for s, _ in out) and "enter each invoice's Amount Paid" in out[0][1][0]
    rows = [_receipt_only(7, amount="1180.00"), _receipt_only(8, amount="500.00")]
    assert {s for s, _ in _evaluate(payments, rows, {"docs/a.pdf": _full_receipt("1680.00")})} == {"READY"}


def test_something_missing_on_both_sheet_and_receipt(payments):
    receipt = _full_receipt()
    del receipt["fields"]["reference_number"]
    [(status, msgs)] = _evaluate(payments, [_receipt_only()], {"docs/a.pdf": receipt})
    assert status == "ERROR" and any("UTR / Reference not in the sheet or on the receipt" in m for m in msgs)


def test_receipt_first_template_requires_mode_and_receipt():
    ws = load_workbook(io.BytesIO(wbk.build_template("PAYMENTS", []))).active
    required = [c.value for c in ws[1] if c.value and c.value.endswith("*")]
    assert required == ["Invoice Number *", "Payment Mode *", "Receipt File Name *"]
    headers = ["Invoice ID", "Invoice Number *", "Payment Mode *", "Receipt File Name *"]
    rows = wbk.parse_workbook("PAYMENTS", _sheet(headers, [7, "INV-7", "NEFT", r"C:\Users\finance\Receipts\advice-7.pdf"]))
    assert rows[0].problems == [] and wbk.file_name_of(rows[0].data["receipt_file"]) == "advice-7.pdf"
    missing_mode = wbk.parse_workbook("PAYMENTS", _sheet(headers, [7, "INV-7", None, "advice-7.pdf"]))
    assert missing_mode[0].problems == ["Payment Mode is required"]
