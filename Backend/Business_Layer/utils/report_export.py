# Backend/Business_Layer/utils/report_export.py
"""Excel / PDF rendering for APReportingService reports.

Both formats render exactly the rows and per-currency totals the JSON report
carries - nothing is recalculated here. Excel via openpyxl and PDF via PyMuPDF,
both already project dependencies (no new packages).
"""
from __future__ import annotations

import datetime
import io
from decimal import Decimal
from typing import Any, Dict, List

import fitz  # PyMuPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BRAND = "0A0082"
_MONEY_FORMAT = "#,##0.00"
_DATE_FORMAT = "DD-MM-YYYY"


def _display(value, kind: str) -> str:
    if value is None or value == "":
        return ""
    if kind == "money":
        return f"{Decimal(str(value)):,.2f}"
    if kind == "date" and isinstance(value, (datetime.date, datetime.datetime)):
        return value.strftime("%d-%m-%Y")
    if kind == "number" and isinstance(value, Decimal):
        return f"{value.normalize():f}"
    return str(value)


def _period_text(report: Dict[str, Any]) -> str:
    period = report.get("period") or {}
    if period.get("type") == "range":
        return f"Period: {_display(period.get('from_date'), 'date')} to {_display(period.get('to_date'), 'date')}"
    return f"As of {_display(period.get('as_of'), 'date')}"


def filename_for(report: Dict[str, Any], extension: str) -> str:
    stamp = datetime.date.today().strftime("%Y%m%d")
    return f"{report['key']}_{stamp}.{extension}"


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def to_xlsx(report: Dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report["key"][:31]
    columns: List[dict] = report["columns"]

    sheet.append([report["title"]])
    sheet["A1"].font = Font(bold=True, size=14, color=BRAND)
    sheet.append([_period_text(report) + f"   ·   Generated {report['generated_at']:%d-%m-%Y %H:%M}"])
    sheet.append([])

    header_row = sheet.max_row + 1
    sheet.append([c["label"] for c in columns])
    for index in range(1, len(columns) + 1):
        cell = sheet.cell(row=header_row, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BRAND)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in report["rows"]:
        values = []
        for c in columns:
            value = row.get(c["key"])
            if c["type"] == "money" and value is not None:
                value = float(value)
            elif c["type"] == "number" and isinstance(value, Decimal):
                value = float(value)
            values.append(value)
        sheet.append(values)
        current = sheet.max_row
        for index, c in enumerate(columns, start=1):
            if c["type"] == "money":
                sheet.cell(row=current, column=index).number_format = _MONEY_FORMAT
            elif c["type"] == "date":
                sheet.cell(row=current, column=index).number_format = _DATE_FORMAT

    last_data_row = sheet.max_row
    if report["rows"]:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{last_data_row}"
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    money_columns = [(i, c) for i, c in enumerate(columns, start=1) if c["type"] == "money"]
    if report.get("totals") and money_columns:
        sheet.append([])
        for total in report["totals"]:
            sheet.append([f"Total ({total['currency_code']}) - {total['rows']} rows"])
            current = sheet.max_row
            sheet.cell(row=current, column=1).font = Font(bold=True)
            for index, c in money_columns:
                cell = sheet.cell(row=current, column=index, value=float(total.get(c["key"], 0)))
                cell.number_format = _MONEY_FORMAT
                cell.font = Font(bold=True)

    for note in report.get("notes") or []:
        sheet.append([])
        sheet.append([note])
        sheet.cell(row=sheet.max_row, column=1).font = Font(italic=True, color="666666")

    for index, c in enumerate(columns, start=1):
        width = max([len(c["label"])] + [len(_display(r.get(c["key"]), c["type"])) for r in report["rows"][:500]])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 48)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_PAGE = fitz.paper_rect("a4-l")
_MARGIN = 28
_ROW_H = 14
_FONT = 7


def _fit(text: str, width: float, font: str = "helv", size: float = _FONT) -> str:
    if fitz.get_text_length(text, fontname=font, fontsize=size) <= width:
        return text
    while text and fitz.get_text_length(text + "…", fontname=font, fontsize=size) > width:
        text = text[:-1]
    return text + "…"


def to_pdf(report: Dict[str, Any]) -> bytes:
    columns: List[dict] = report["columns"]
    usable = _PAGE.width - 2 * _MARGIN
    weights = [1.6 if c["type"] == "text" else 1.0 for c in columns]
    widths = [usable * w / sum(weights) for w in weights]

    doc = fitz.open()
    state = {"page": None, "y": 0}

    def new_page():
        page = doc.new_page(width=_PAGE.width, height=_PAGE.height)
        page.insert_text((_MARGIN, _MARGIN + 10), report["title"], fontsize=12, fontname="hebo", color=(0.04, 0, 0.51))
        page.insert_text((_MARGIN, _MARGIN + 24), _period_text(report) + f"   ·   Generated {report['generated_at']:%d-%m-%Y %H:%M}",
                         fontsize=7.5, fontname="helv", color=(0.35, 0.35, 0.4))
        y = _MARGIN + 40
        page.draw_rect(fitz.Rect(_MARGIN, y - 10, _PAGE.width - _MARGIN, y + 4), color=None, fill=(0.04, 0, 0.51))
        x = _MARGIN
        for c, w in zip(columns, widths):
            label = _fit(c["label"], w - 4, "hebo")
            if c["type"] in ("money", "int", "number"):
                page.insert_text((x + w - 2 - fitz.get_text_length(label, fontname="hebo", fontsize=_FONT), y),
                                 label, fontsize=_FONT, fontname="hebo", color=(1, 1, 1))
            else:
                page.insert_text((x + 2, y), label, fontsize=_FONT, fontname="hebo", color=(1, 1, 1))
            x += w
        state["page"], state["y"] = page, y + _ROW_H

    def write_row(values, bold=False, shade=False):
        if state["page"] is None or state["y"] > _PAGE.height - _MARGIN - 20:
            new_page()
        page, y = state["page"], state["y"]
        font = "hebo" if bold else "helv"
        if shade:
            page.draw_rect(fitz.Rect(_MARGIN, y - 9, _PAGE.width - _MARGIN, y + 4), color=None, fill=(0.95, 0.95, 0.98))
        x = _MARGIN
        for (c, w), value in zip(zip(columns, widths), values):
            text = _fit(value, w - 4, font)
            if c["type"] in ("money", "int", "number"):
                page.insert_text((x + w - 2 - fitz.get_text_length(text, fontname=font, fontsize=_FONT), y),
                                 text, fontsize=_FONT, fontname=font)
            else:
                page.insert_text((x + 2, y), text, fontsize=_FONT, fontname=font)
            x += w
        state["y"] = y + _ROW_H

    if not report["rows"]:
        new_page()
        state["page"].insert_text((_MARGIN, state["y"] + 6), "No rows for the selected filters.", fontsize=9, fontname="helv")
    for index, row in enumerate(report["rows"]):
        write_row([_display(row.get(c["key"]), c["type"]) for c in columns], shade=index % 2 == 1)

    money_keys = {c["key"] for c in columns if c["type"] == "money"}
    for total in report.get("totals") or []:
        if not money_keys:
            break
        values = []
        for i, c in enumerate(columns):
            if c["key"] in money_keys:
                values.append(_display(total.get(c["key"]), "money"))
            else:
                values.append(f"Total {total['currency_code']} ({total['rows']})" if i == 0 else "")
        write_row(values, bold=True)

    for note in report.get("notes") or []:
        if state["y"] > _PAGE.height - _MARGIN - 20:
            new_page()
        state["page"].insert_text((_MARGIN, state["y"] + 6), _fit(note, usable, size=7), fontsize=7, fontname="helv",
                                  color=(0.4, 0.4, 0.45))
        state["y"] += 12

    for number, page in enumerate(doc, start=1):
        page.insert_text((_PAGE.width - _MARGIN - 50, _PAGE.height - 14), f"Page {number} of {len(doc)}",
                         fontsize=7, fontname="helv", color=(0.4, 0.4, 0.45))
    return doc.tobytes(garbage=3, deflate=True)
